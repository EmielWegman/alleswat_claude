#!/usr/bin/env python3
"""Parse _vragen_bron.txt (delimiter ~|~) naar CSV en XLSX.

Kolommen: Categorie | Onderdeel | Nr | Paragraaf | Vraag
Bron: ActiZ 'Voorbeeldvragen ter verduidelijking en beinvloeding Wlz inkoopbeleid 2027-2029'
Scope: landelijke vragen (I, II, III) + Zorgkantoor Zilveren Kruis.
"""
import csv
import sys
from pathlib import Path

DELIM = "~|~"
BASE = Path(__file__).resolve().parent
SRC = BASE / "_vragen_bron.txt"
CSV_OUT = BASE / "vragen-landelijk-en-zilveren-kruis-2027.csv"
XLSX_OUT = BASE / "vragen-landelijk-en-zilveren-kruis-2027.xlsx"
HEADER = ["Categorie", "Onderdeel", "Nr", "Paragraaf", "Vraag"]

def parse(path):
    rows = []
    for lineno, raw in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
        if not raw.strip():
            continue
        parts = raw.split(DELIM)
        if len(parts) != 5:
            sys.exit(f"FOUT regel {lineno}: {len(parts)} velden i.p.v. 5 -> {raw[:80]!r}")
        cat, onderdeel, nr, para, vraag = (p.strip() for p in parts)
        if not vraag.endswith(("?", ".")):
            print(f"  let op regel {lineno}: vraag eindigt niet op ? of . -> ...{vraag[-40:]!r}")
        rows.append([cat, onderdeel, nr, para, vraag])
    return rows

def write_csv(rows):
    # UTF-8 met BOM zodat Excel de accenten direct goed toont
    with CSV_OUT.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";", quoting=csv.QUOTE_ALL)
        w.writerow(HEADER)
        w.writerows(rows)

def write_xlsx(rows):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("openpyxl niet beschikbaar - alleen CSV geschreven.")
        return False
    wb = Workbook()
    ws = wb.active
    ws.title = "Vragen 2027"
    ws.append(HEADER)
    for r in rows:
        ws.append(r)
    # opmaak header
    hfill = PatternFill("solid", fgColor="4F2D7F")
    hfont = Font(bold=True, color="FFFFFF")
    for c in range(1, len(HEADER) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = hfill
        cell.font = hfont
        cell.alignment = Alignment(vertical="center")
    widths = {"A": 16, "B": 38, "C": 6, "D": 22, "E": 100}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        row[4].alignment = Alignment(wrap_text=True, vertical="top")
        row[2].alignment = Alignment(horizontal="center", vertical="top")
        for cell in (row[0], row[1], row[3]):
            cell.alignment = Alignment(vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}{len(rows)+1}"
    wb.save(XLSX_OUT)
    return True

def main():
    rows = parse(SRC)
    write_csv(rows)
    has_xlsx = write_xlsx(rows)
    # samenvatting
    from collections import Counter
    by_part = Counter(r[1] for r in rows)
    print(f"\nTotaal vragen: {len(rows)}")
    for part, n in by_part.items():
        print(f"  {part}: {n}")
    print(f"\nCSV : {CSV_OUT.name}")
    if has_xlsx:
        print(f"XLSX: {XLSX_OUT.name}")

if __name__ == "__main__":
    main()
